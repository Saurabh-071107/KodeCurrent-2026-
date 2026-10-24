"""
Unit and Integration Tests for DocMultilingualModel Pipeline.
Tests:
- Text extraction from DOCX, PDF, and TXT
- Intelligent chunking
- Translation to Hindi, Marathi, Gujarati, English, Telugu
- Exporting to DOCX, PDF, and TXT
- Voiceover Audio Synthesis (MP3) for all 5 languages
"""

import os
import unittest
import docx
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph
from reportlab.lib.styles import getSampleStyleSheet

from doc_translator import (
    DocMultilingualModel,
    DocumentExtractor,
    TranslationEngine,
    DocumentExporter,
    VoiceoverGenerator,
    SUPPORTED_LANGUAGES,
)


class TestMultilingualTranslator(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.test_dir = os.path.join(os.path.dirname(__file__), "test_files")
        os.makedirs(cls.test_dir, exist_ok=True)

        cls.sample_text = (
            "Natural language processing enables computers to understand human language.\n\n"
            "India is a diverse country with many spoken languages such as Hindi, Marathi, Gujarati, and Telugu.\n\n"
            "Technology bridges communication gaps and connects people worldwide."
        )

        # 1. Create Sample TXT file
        cls.txt_path = os.path.join(cls.test_dir, "sample.txt")
        with open(cls.txt_path, "w", encoding="utf-8") as f:
            f.write(cls.sample_text)

        # 2. Create Sample DOCX file
        cls.docx_path = os.path.join(cls.test_dir, "sample.docx")
        doc = docx.Document()
        doc.add_heading("NLP and Multilingual Systems", level=1)
        for p in cls.sample_text.split("\n\n"):
            doc.add_paragraph(p)
        doc.save(cls.docx_path)

        # 3. Create Sample PDF file
        cls.pdf_path = os.path.join(cls.test_dir, "sample.pdf")
        pdf_doc = SimpleDocTemplate(cls.pdf_path, pagesize=letter)
        styles = getSampleStyleSheet()
        story = [Paragraph("NLP and Multilingual Systems", styles["Title"])]
        for p in cls.sample_text.split("\n\n"):
            story.append(Paragraph(p, styles["Normal"]))
        pdf_doc.build(story)

    def test_supported_languages(self):
        """Ensure all 5 requested languages are properly registered."""
        expected = {"hi", "mr", "gu", "en", "te"}
        self.assertTrue(expected.issubset(set(SUPPORTED_LANGUAGES.keys())))
        for lang in expected:
            self.assertIn("name", SUPPORTED_LANGUAGES[lang])
            self.assertIn("native", SUPPORTED_LANGUAGES[lang])

    def test_extract_txt(self):
        """Verify text extraction from plain text file."""
        data = DocumentExtractor.extract(self.txt_path)
        self.assertEqual(data["format"], "txt")
        self.assertGreater(len(data["paragraphs"]), 0)
        self.assertIn("Natural language processing", data["full_text"])

    def test_extract_docx(self):
        """Verify text extraction from DOCX file."""
        data = DocumentExtractor.extract(self.docx_path)
        self.assertEqual(data["format"], "docx")
        self.assertGreaterEqual(len(data["paragraphs"]), 3)
        self.assertIn("India is a diverse country", data["full_text"])

    def test_extract_pdf(self):
        """Verify text extraction from PDF file."""
        data = DocumentExtractor.extract(self.pdf_path)
        self.assertEqual(data["format"], "pdf")
        self.assertGreater(len(data["paragraphs"]), 0)
        self.assertIn("Natural language processing", data["full_text"])

    def test_chunking(self):
        """Verify intelligent chunking preserves paragraph structures."""
        paragraphs = ["Short paragraph 1.", "Short paragraph 2.", "Short paragraph 3."]
        chunks = DocumentExtractor.create_translation_chunks(paragraphs, max_chunk_chars=100)
        self.assertGreater(len(chunks), 0)

    def test_translation_engine_single(self):
        """Verify direct single text translation for all 5 languages."""
        engine = TranslationEngine()
        sample = "Good morning! Welcome to the new world."

        for lang_code in ["hi", "mr", "gu", "en", "te"]:
            translated = engine.translate_text(sample, target_lang=lang_code)
            self.assertIsInstance(translated, str)
            self.assertGreater(len(translated), 0)

    def test_voiceover_generation_all_five_languages(self):
        """Verify MP3 voiceover generation for Hindi, Marathi, Gujarati, English, and Telugu."""
        samples = {
            "hi": "नमस्ते, यह हिंदी में एक परीक्षण है।",
            "mr": "नमस्कार, ही मराठीतील एक चाचणी आहे.",
            "gu": "નમસ્તે, આ ગુજરાતીમાં એક પરીક્ષણ છે.",
            "en": "Hello, this is a test in English.",
            "te": "నమస్కారం, ఇది తెలుగులో ఒక పరీక్ష.",
        }

        for lang, phrase in samples.items():
            audio_bytes = VoiceoverGenerator.generate_audio_bytes(phrase, lang)
            self.assertIsInstance(audio_bytes, bytes)
            self.assertGreater(len(audio_bytes), 1000, f"Voiceover for {lang} is too small")

    def test_full_model_pipeline_with_voiceovers(self):
        """Verify end-to-end model execution on DOCX with translations and voiceover exports."""
        model = DocMultilingualModel()
        result = model.translate_document(
            self.docx_path,
            target_langs=["hi", "mr", "gu", "en", "te"],
            parallel=True,
        )

        self.assertEqual(result.format, "docx")
        self.assertEqual(len(result.translations), 5)

        for lang in ["hi", "mr", "gu", "en", "te"]:
            text = result.get_text(lang)
            self.assertGreater(len(text), 10, f"Translation for {lang} is empty")

        # Test export to documents
        docx_bytes = result.to_bytes("hi", "docx")
        self.assertGreater(len(docx_bytes), 100)

        pdf_bytes = result.to_bytes("gu", "pdf")
        self.assertGreater(len(pdf_bytes), 100)

        txt_bytes = result.to_bytes("te", "txt")
        self.assertGreater(len(txt_bytes), 10)

        # Test voiceover audio generation from result
        hi_audio = result.to_audio_bytes("hi")
        self.assertGreater(len(hi_audio), 1000)

        # Test bulk save including MP3 voiceovers
        out_dir = os.path.join(self.test_dir, "output_with_audio")
        saved_files = result.save_all(out_dir, formats=["docx", "txt", "mp3"])
        self.assertEqual(len(saved_files), 5)
        for lang, files in saved_files.items():
            mp3_files = [f for f in files if f.endswith(".mp3")]
            self.assertEqual(len(mp3_files), 1, f"Missing MP3 for {lang}")
            self.assertTrue(os.path.exists(mp3_files[0]))


if __name__ == "__main__":
    unittest.main()
