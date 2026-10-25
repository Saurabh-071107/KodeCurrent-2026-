"""Document parsing module supporting PDF, DOCX, and TXT."""
import os
import re
from pathlib import Path
from typing import List, Dict, Any, Optional, Tuple
from dataclasses import dataclass
import structlog

logger = structlog.get_logger(__name__)


@dataclass
class ParsedPage:
    """A single page from a document."""
    page_number: int
    text: str
    metadata: Dict[str, Any]


@dataclass
class ParsedDocument:
    """Parsed document with pages and metadata."""
    filename: str
    file_type: str
    pages: List[ParsedPage]
    total_pages: int
    metadata: Dict[str, Any]


class DocumentParsingError(Exception):
    """Document parsing specific errors."""
    pass


class UnsupportedFileTypeError(DocumentParsingError):
    """Raised when file type is not supported."""
    pass


class FileTooLargeError(DocumentParsingError):
    """Raised when file exceeds size limit."""
    pass


class CorruptedFileError(DocumentParsingError):
    """Raised when file is corrupted or unreadable."""
    pass


def validate_file(file_path: str | Path, max_size_mb: int, allowed_extensions: List[str]) -> Path:
    """Validate file before parsing. Returns Path object."""
    file_path = Path(file_path)
    if not file_path.exists():
        raise DocumentParsingError(f"File not found: {file_path}")

    # Check file size
    file_size_mb = file_path.stat().st_size / (1024 * 1024)
    if file_size_mb > max_size_mb:
        raise FileTooLargeError(
            f"File size ({file_size_mb:.1f} MB) exceeds limit ({max_size_mb} MB)"
        )

    # Check extension
    ext = file_path.suffix.lower()
    if ext not in allowed_extensions:
        raise UnsupportedFileTypeError(
            f"Unsupported file type: {ext}. Allowed: {allowed_extensions}"
        )
    
    return file_path


def parse_pdf(file_path: Path) -> ParsedDocument:
    """Parse PDF using pypdf."""
    try:
        from pypdf import PdfReader
    except ImportError:
        raise DocumentParsingError("pypdf not installed. Install with: pip install pypdf")

    pages = []
    metadata = {}

    try:
        reader = PdfReader(file_path)
        pdf_metadata = reader.metadata or {}
        metadata = {
            "page_count": len(reader.pages),
            "title": pdf_metadata.get("/Title", "") or "",
            "author": pdf_metadata.get("/Author", "") or "",
            "subject": pdf_metadata.get("/Subject", "") or "",
            "creator": pdf_metadata.get("/Creator", "") or "",
            "producer": pdf_metadata.get("/Producer", "") or "",
            "creation_date": str(pdf_metadata.get("/CreationDate", "")) or "",
            "modification_date": str(pdf_metadata.get("/ModDate", "")) or "",
        }

        for page_num, page in enumerate(reader.pages):
            try:
                text = page.extract_text()
            except Exception:
                text = ""
            
            page_metadata = {
                "rotation": getattr(page, "rotation", 0),
            }
            pages.append(ParsedPage(
                page_number=page_num + 1,
                text=text,
                metadata=page_metadata,
            ))

    except Exception as e:
        if "corrupt" in str(e).lower() or "damaged" in str(e).lower():
            raise CorruptedFileError(f"PDF appears to be corrupted: {e}")
        raise DocumentParsingError(f"Failed to parse PDF: {e}")

    if not pages:
        raise DocumentParsingError("PDF contains no extractable pages")

    logger.info("pdf_parsed", filename=file_path.name, pages=len(pages))
    return ParsedDocument(
        filename=file_path.name,
        file_type="pdf",
        pages=pages,
        total_pages=len(pages),
        metadata=metadata,
    )


def parse_docx(file_path: Path) -> ParsedDocument:
    """Parse DOCX using python-docx."""
    try:
        from docx import Document
    except ImportError:
        raise DocumentParsingError("python-docx not installed. Install with: pip install python-docx")

    pages = []
    metadata = {}

    try:
        doc = Document(file_path)
        metadata = {
            "paragraph_count": len(doc.paragraphs),
            "table_count": len(doc.tables),
            "core_properties": {
                "title": doc.core_properties.title or "",
                "author": doc.core_properties.author or "",
                "subject": doc.core_properties.subject or "",
                "created": str(doc.core_properties.created) if doc.core_properties.created else "",
                "modified": str(doc.core_properties.modified) if doc.core_properties.modified else "",
            },
        }

        # Extract text from paragraphs
        all_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                all_text.append(para.text)

        # Extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text for cell in row.cells)
                if row_text.strip():
                    all_text.append(row_text)

        # Combine into a single "page" for DOCX (no native page concept)
        text = "\n\n".join(all_text)
        pages.append(ParsedPage(
            page_number=1,
            text=text,
            metadata={"paragraph_count": len(doc.paragraphs)},
        ))

    except Exception as e:
        if "corrupt" in str(e).lower() or "damaged" in str(e).lower():
            raise CorruptedFileError(f"DOCX appears to be corrupted: {e}")
        raise DocumentParsingError(f"Failed to parse DOCX: {e}")

    if not pages or not pages[0].text.strip():
        raise DocumentParsingError("DOCX contains no extractable text")

    logger.info("docx_parsed", filename=file_path.name, paragraphs=len(doc.paragraphs))
    return ParsedDocument(
        filename=file_path.name,
        file_type="docx",
        pages=pages,
        total_pages=1,
        metadata=metadata,
    )


def parse_txt(file_path: Path) -> ParsedDocument:
    """Parse plain text file."""
    try:
        # Try to detect encoding
        import chardet
        with open(file_path, "rb") as f:
            raw = f.read()
        detected = chardet.detect(raw)
        encoding = detected.get("encoding", "utf-8")
    except ImportError:
        encoding = "utf-8"

    try:
        text = file_path.read_text(encoding=encoding)
    except UnicodeDecodeError:
        # Fallback: try common encodings
        for enc in ["utf-8", "latin-1", "cp1252", "iso-8859-1"]:
            try:
                text = file_path.read_text(encoding=enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise CorruptedFileError("Unable to decode text file with any common encoding")

    if not text.strip():
        raise DocumentParsingError("Text file is empty")

    pages = [ParsedPage(
        page_number=1,
        text=text,
        metadata={"encoding": encoding, "char_count": len(text)},
    )]

    metadata = {"encoding": encoding, "char_count": len(text)}
    logger.info("txt_parsed", filename=file_path.name, chars=len(text))
    return ParsedDocument(
        filename=file_path.name,
        file_type="txt",
        pages=pages,
        total_pages=1,
        metadata=metadata,
    )


def parse_document(file_path: Path, max_size_mb: int = 50, allowed_extensions: List[str] = None) -> ParsedDocument:
    """
    Parse a document file based on its extension.
    
    Args:
        file_path: Path to the document file
        max_size_mb: Maximum file size in MB
        allowed_extensions: List of allowed file extensions
        
    Returns:
        ParsedDocument with pages and metadata
        
    Raises:
        DocumentParsingError: If parsing fails
        UnsupportedFileTypeError: If file type not supported
        FileTooLargeError: If file exceeds size limit
    """
    if allowed_extensions is None:
        allowed_extensions = [".pdf", ".docx", ".txt"]

    file_path = validate_file(file_path, max_size_mb, allowed_extensions)

    ext = file_path.suffix.lower()

    if ext == ".pdf":
        return parse_pdf(file_path)
    elif ext == ".docx":
        return parse_docx(file_path)
    elif ext == ".txt":
        return parse_txt(file_path)
    else:
        raise UnsupportedFileTypeError(f"Unsupported file type: {ext}")


def extract_text_for_chunking(parsed_doc: ParsedDocument) -> List[Tuple[str, Dict[str, Any]]]:
    """
    Extract text chunks with page/paragraph references for chunking.
    
    Returns:
        List of (text, metadata) tuples where metadata contains page_number, etc.
    """
    chunks = []
    for page in parsed_doc.pages:
        if page.text.strip():
            chunks.append((
                page.text,
                {
                    "page_number": page.page_number,
                    "source_filename": parsed_doc.filename,
                    "file_type": parsed_doc.file_type,
                    **page.metadata,
                },
            ))
    return chunks


def convert_to_markdown(file_path: Path) -> str:
    """
    Convert any supported document (PDF, DOCX, TXT, etc.) to structured Markdown using Microsoft MarkItDown.
    Preserves tables, headers, statutory sections, and legal schedules.
    """
    try:
        from markitdown import MarkItDown
        md = MarkItDown()
        result = md.convert(str(file_path))
        if result and result.text_content:
            logger.info("markitdown_conversion_success", filename=file_path.name, chars=len(result.text_content))
            return result.text_content
    except Exception as e:
        logger.warning("markitdown_conversion_fallback", error=str(e), filename=file_path.name)

    # Fallback to standard parser
    parsed = parse_document(file_path)
    return "\n\n".join(page.text for page in parsed.pages if page.text.strip())