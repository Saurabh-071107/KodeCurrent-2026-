"""Document export feature module."""
import os
import uuid
import shutil
import structlog
from typing import Optional
from pathlib import Path
from datetime import datetime, timedelta

from app.config import get_settings
from app.schemas import (
    ExportRequest,
    ExportResponse,
    ExportFormat,
    DraftResponse,
    DraftSection,
)

logger = structlog.get_logger(__name__)


class ExportError(Exception):
    """Export specific errors."""
    pass


async def export_document(request: ExportRequest) -> ExportResponse:
    """
    Export a draft document to DOCX or PDF format.
    
    Args:
        request: ExportRequest with draft and format options
        
    Returns:
        ExportResponse with download information
    """
    settings = get_settings()
    export_dir = Path(settings.export.output_dir)
    export_dir.mkdir(parents=True, exist_ok=True)
    
    # Generate unique file ID and filename
    file_id = str(uuid.uuid4())
    timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    
    if request.filename:
        base_name = request.filename
    else:
        base_name = f"{request.draft.document_type.value}_{timestamp}"
    
    # Generate DOCX first (always)
    docx_filename = f"{file_id}_{base_name}.docx"
    docx_path = export_dir / docx_filename
    
    try:
        _generate_docx(request.draft, docx_path, request.include_comments, request.include_changes_summary)
    except Exception as e:
        logger.error("docx_generation_failed", error=str(e))
        raise ExportError(f"DOCX generation failed: {str(e)}")
    
    final_path = docx_path
    final_format = ExportFormat.DOCX
    final_filename = docx_filename
    
    # Convert to PDF if requested
    if request.format == ExportFormat.PDF:
        pdf_filename = f"{file_id}_{base_name}.pdf"
        pdf_path = export_dir / pdf_filename
        
        try:
            _convert_to_pdf(docx_path, pdf_path, settings.export.libreoffice_path)
            final_path = pdf_path
            final_format = ExportFormat.PDF
            final_filename = pdf_filename
        except Exception as e:
            logger.warning("pdf_conversion_failed", error=str(e))
            # If PDF conversion fails, return DOCX with warning
            logger.warning("pdf_conversion_failed_returning_docx", error=str(e))
            # We'll return the DOCX but note the format
            pass
    
    # Get file size
    file_size = final_path.stat().st_size
    
    # Calculate expiration
    expires_at = datetime.utcnow() + timedelta(hours=settings.export.max_file_age_hours)
    
    # Generate download URL (in production, this would be a signed URL or served via a file endpoint)
    download_url = f"/exports/download/{final_filename}"
    
    logger.info("export_completed", file_id=file_id, format=final_format.value, size=file_size)
    
    return ExportResponse(
        file_id=file_id,
        filename=final_filename,
        format=final_format,
        download_url=download_url,
        expires_at=expires_at,
        file_size_bytes=file_size,
    )


def _generate_docx(
    draft: DraftResponse,
    output_path: Path,
    include_comments: bool = False,
    include_changes_summary: bool = False,
) -> None:
    """Generate a DOCX file from a draft response."""
    from docx import Document
    from docx.shared import Inches, Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15
    
    # Title
    title_para = doc.add_heading(draft.title, level=0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_para.runs[0].font.size = Pt(16)
    title_para.runs[0].font.bold = True
    
    doc.add_paragraph()  # spacing
    
    # Add sections
    for section in draft.sections:
        if section.title:
            # Add section heading
            heading = doc.add_heading(section.title, level=1)
            heading.runs[0].font.size = Pt(14)
            heading.runs[0].font.bold = True
        
        # Add section content
        if section.content:
            # Split content into paragraphs
            paragraphs = section.content.split('\n\n')
            for para_text in paragraphs:
                para_text = para_text.strip()
                if not para_text:
                    continue
                
                # Check for list items
                lines = para_text.split('\n')
                if len(lines) > 1 and all(line.strip().startswith(('•', '-', '*', '1.', '2.', '3.', 'a.', 'b.', 'c.', 'i.', 'ii.', 'iii.')) for line in lines if line.strip()):
                    # This looks like a list
                    for line in lines:
                        line = line.strip()
                        if line:
                            p = doc.add_paragraph(line, style='List Bullet')
                            p.paragraph_format.space_after = Pt(3)
                else:
                    p = doc.add_paragraph(para_text)
                    p.paragraph_format.space_after = Pt(6)
        
        doc.add_paragraph()  # spacing between sections
    
    # Add disclaimer if present
    if draft.disclaimer:
        doc.add_paragraph()
        disclaimer_para = doc.add_paragraph(draft.disclaimer)
        disclaimer_para.runs[0].font.size = Pt(9)
        disclaimer_para.runs[0].font.italic = True
        disclaimer_para.runs[0].font.color.rgb = RGBColor(0x80, 0x80, 0x80)
        disclaimer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Save
    doc.save(output_path)


def _convert_to_pdf(docx_path: Path, pdf_path: Path, libreoffice_path: Optional[str] = None) -> None:
    """Convert DOCX to PDF using LibreOffice."""
    import subprocess
    
    # Find LibreOffice
    if libreoffice_path and Path(libreoffice_path).exists():
        soffice = libreoffice_path
    else:
        # Try to find LibreOffice automatically
        possible_paths = [
            # Windows
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
            # Linux
            "/usr/bin/libreoffice",
            "/usr/bin/soffice",
            # macOS
            "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        ]
        
        soffice = None
        for path in possible_paths:
            if Path(path).exists():
                soffice = path
                break
        
        if not soffice:
            raise ExportError(
                "LibreOffice not found. Install LibreOffice or set LIBREOFFICE_PATH environment variable. "
                "PDF export requires LibreOffice for DOCX to PDF conversion."
            )
    
    # Convert using LibreOffice headless
    output_dir = pdf_path.parent
    cmd = [
        soffice,
        "--headless",
        "--convert-to", "pdf",
        "--outdir", str(output_dir),
        str(docx_path),
    ]
    
    try:
        result = subprocess.run(cmd, capture_output=True, text=True, timeout=120)
        if result.returncode != 0:
            raise ExportError(f"LibreOffice conversion failed: {result.stderr}")
        
        # LibreOffice creates PDF with same name as input
        expected_pdf = output_dir / (docx_path.stem + ".pdf")
        if expected_pdf.exists():
            if expected_pdf != pdf_path:
                shutil.move(str(expected_pdf), str(pdf_path))
        else:
            raise ExportError("PDF file not created by LibreOffice")
            
    except subprocess.TimeoutExpired:
        raise ExportError("LibreOffice conversion timed out")
    except FileNotFoundError:
        raise ExportError("LibreOffice executable not found")
    except Exception as e:
        raise ExportError(f"PDF conversion failed: {str(e)}")


async def cleanup_old_exports() -> int:
    """Clean up export files older than max_file_age_hours."""
    settings = get_settings()
    export_dir = Path(settings.export.output_dir)
    
    if not export_dir.exists():
        return 0
    
    cutoff = datetime.utcnow() - timedelta(hours=settings.export.max_file_age_hours)
    deleted_count = 0
    
    for file_path in export_dir.iterdir():
        if file_path.is_file():
            mtime = datetime.fromtimestamp(file_path.stat().st_mtime)
            if mtime < cutoff:
                try:
                    file_path.unlink()
                    deleted_count += 1
                except Exception as e:
                    logger.warning("export_cleanup_failed", file=str(file_path), error=str(e))
    
    if deleted_count > 0:
        logger.info("export_cleanup_completed", deleted_count=deleted_count)
    
    return deleted_count