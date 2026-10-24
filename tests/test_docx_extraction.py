"""Tests for DOCX extraction, structure parsing, and metadata preservation."""

import io
import pytest
import docx

from draftmate_ai.document_processing.docx_extractor import DOCXExtractor
from draftmate_ai.shared.document_models import ElementType
from draftmate_ai.shared.exceptions import EmptyDocumentError, FileCorruptedError


def test_docx_extraction_valid(sample_docx_bytes):
    """Test extracting structure, headings, clauses, and tables from DOCX."""
    extractor = DOCXExtractor()
    doc = extractor.extract(sample_docx_bytes, filename="nda.docx")

    assert doc.metadata.filename == "nda.docx"
    assert doc.metadata.file_type == "docx"
    assert doc.metadata.title == "Non-Disclosure Agreement"
    assert doc.metadata.author == "DraftMate Legal Team"
    assert doc.metadata.page_count >= 1

    # Verify headings
    headings = doc.get_headings()
    assert len(headings) >= 2
    heading_texts = [h.normalized_text for h in headings]
    assert any("NON-DISCLOSURE" in h for h in heading_texts)
    assert any("CONFIDENTIALITY" in h for h in heading_texts)

    # Verify legal clauses
    clauses = doc.get_clauses()
    assert len(clauses) >= 2
    clause_numbers = [c.clause_number for c in clauses if c.clause_number]
    assert any("1.1" in cn or "Section 1.1" in cn for cn in clause_numbers)
    assert any("2.1" in cn or "Section 2.1" in cn for cn in clause_numbers)

    # Verify table extraction
    tables = doc.get_tables()
    assert len(tables) == 1
    table = tables[0]
    assert "Party Role" in table.headers
    assert "Designated Representative" in table.headers
    assert len(table.rows) == 2
    assert table.rows[0][0] == "Discloser"
    assert table.rows[0][1] == "Alice Smith"
    assert "| Party Role | Designated Representative |" in table.to_markdown()


def test_docx_extraction_empty_bytes():
    """Test that zero-byte DOCX raises EmptyDocumentError."""
    extractor = DOCXExtractor()
    with pytest.raises(EmptyDocumentError):
        extractor.extract(b"", filename="empty.docx")


def test_docx_extraction_corrupted_header():
    """Test that non-zip DOCX raises FileCorruptedError."""
    extractor = DOCXExtractor()
    with pytest.raises(FileCorruptedError):
        extractor.extract(b"Not a zip file header at all", filename="invalid.docx")


def test_docx_page_break_tracking():
    """Test tracking page boundaries when explicit page breaks are added."""
    doc = docx.Document()
    doc.add_heading("PAGE ONE HEADER", level=1)
    doc.add_paragraph("Content on page one.")
    doc.add_page_break()
    doc.add_heading("PAGE TWO HEADER", level=1)
    doc.add_paragraph("Content on page two.")

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    extractor = DOCXExtractor()
    processed = extractor.extract(docx_bytes, filename="multipage.docx")

    assert processed.metadata.page_count == 2
    p1_elems = processed.get_elements_by_page(1)
    p2_elems = processed.get_elements_by_page(2)
    assert len(p1_elems) >= 2
    assert len(p2_elems) >= 2

