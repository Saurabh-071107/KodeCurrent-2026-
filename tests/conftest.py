"""Shared pytest fixtures creating synthetic test documents for PDF, DOCX, and TXT."""

import io
import pytest
from pathlib import Path
from PIL import Image, ImageDraw

import pymupdf
import docx


@pytest.fixture
def sample_pdf_bytes() -> bytes:
    """Generate a clean synthetic PDF with headings, paragraphs, clauses, and tables."""
    doc = pymupdf.open()

    # Page 1
    page1 = doc.new_page(width=595, height=842)
    page1.insert_text((50, 60), "CONFIDENTIAL MASTER SERVICES AGREEMENT", fontsize=18, fontname="helv")
    page1.insert_text((50, 100), "This Master Services Agreement is entered into by and between Alpha Corp and Beta LLC.", fontsize=11)
    page1.insert_text((50, 130), "ARTICLE I - DEFINITIONS", fontsize=14, fontname="helv")
    page1.insert_text((50, 160), "Section 1.01 Defined Terms. For the purposes of this Agreement, the following terms have meanings set forth below.", fontsize=11)
    page1.insert_text((50, 190), "(a) Confidential Information means all proprietary technical, business, and financial data.", fontsize=11)
    page1.insert_text((50, 220), "(b) Effective Date means October 25, 2026.", fontsize=11)
    page1.insert_text((50, 260), "ARTICLE II - OBLIGATIONS", fontsize=14, fontname="helv")
    page1.insert_text((50, 290), "Section 2.01 Standard of Care. The Contractor shall perform all services with utmost professional diligence.", fontsize=11)

    # Page 2
    page2 = doc.new_page(width=595, height=842)
    page2.insert_text((50, 60), "ARTICLE III - FEES AND COMPENSATION", fontsize=14, fontname="helv")
    page2.insert_text((50, 90), "Section 3.01 Invoicing. Invoices shall be rendered on a monthly basis and paid within 30 days.", fontsize=11)
    page2.insert_text((50, 130), "Section 3.02 Fee Schedule.", fontsize=12)

    # Insert a simple drawing or text table on Page 2
    page2.insert_text((50, 160), "Milestone 1 | Initial Architecture | $15,000", fontsize=10)
    page2.insert_text((50, 180), "Milestone 2 | Implementation       | $25,000", fontsize=10)
    page2.insert_text((50, 200), "Milestone 3 | Verification & UAT   | $10,000", fontsize=10)

    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


@pytest.fixture
def sample_scanned_pdf_bytes() -> bytes:
    """Generate a synthetic scanned PDF containing only an image (no extractable text)."""
    # Create an image containing simulated scanned text
    img = Image.new("RGB", (600, 400), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.text((40, 50), "EXHIBIT A - SCANNED SIGNATURE PAGE", fill=(0, 0, 0))
    draw.text((40, 90), "IN WITNESS WHEREOF, the parties hereto have executed this Agreement.", fill=(0, 0, 0))
    draw.text((40, 130), "Signature: John Doe, Chief Executive Officer", fill=(0, 0, 0))

    img_buffer = io.BytesIO()
    img.save(img_buffer, format="PNG")
    img_bytes = img_buffer.getvalue()

    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    rect = pymupdf.Rect(50, 50, 545, 400)
    page.insert_image(rect, stream=img_bytes)

    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


@pytest.fixture
def sample_docx_bytes() -> bytes:
    """Generate a clean synthetic DOCX with headings, paragraphs, clauses, and tables."""
    doc = docx.Document()
    doc.core_properties.title = "Non-Disclosure Agreement"
    doc.core_properties.author = "DraftMate Legal Team"

    doc.add_heading("MUTUAL NON-DISCLOSURE AGREEMENT", level=1)
    doc.add_paragraph("This Agreement is made between Disclosing Party and Receiving Party.")

    doc.add_heading("ARTICLE 1: CONFIDENTIALITY", level=2)
    doc.add_paragraph("Section 1.1 Scope. Receiving Party agrees to hold all Information in strict confidence.")
    doc.add_paragraph("Section 1.2 Exceptions. Confidential Information shall not include publicly known data.")

    doc.add_heading("ARTICLE 2: TERM AND TERMINATION", level=2)
    doc.add_paragraph("Section 2.1 Term. This Agreement remains valid for a period of three (3) years from execution.")

    # Add a table
    table = doc.add_table(rows=3, cols=2)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Party Role"
    hdr_cells[1].text = "Designated Representative"

    row1 = table.rows[1].cells
    row1[0].text = "Discloser"
    row1[1].text = "Alice Smith"

    row2 = table.rows[2].cells
    row2[0].text = "Recipient"
    row2[1].text = "Bob Jones"

    docx_buffer = io.BytesIO()
    doc.save(docx_buffer)
    return docx_buffer.getvalue()


@pytest.fixture
def sample_legal_text() -> str:
    """Sample plain legal text document."""
    return """COMMERCIAL LEASE AGREEMENT

This Commercial Lease Agreement (the "Lease") is made effective as of January 1, 2026.

1. DEMISE OF PREMISES
Landlord hereby leases to Tenant the commercial property located at 100 Main Street, Suite 400.

2. RENT AND PAYMENTS
Section 2.1 Base Rent. Tenant shall pay Landlord base rent in the amount of $5,000 per month.
Section 2.2 Late Charge. Any payment not received within five (5) days of due date incurs a 5% fee.

3. USE OF PREMISES
Tenant shall use the premises strictly for general office and administrative purposes.
(a) Tenant shall comply with all applicable local, state, and federal zoning regulations.
(b) Tenant shall not create a nuisance or disturb other building occupants.

4. INDEMNIFICATION AND INSURANCE
Tenant agrees to indemnify, defend, and hold harmless Landlord against all claims and damages.
"""

