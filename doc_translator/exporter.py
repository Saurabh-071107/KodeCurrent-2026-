"""
Document Exporter module.
Exports translated content to DOCX, PDF, TXT, and JSON formats.
Includes full font styling support for Hindi, Marathi, Gujarati, English, and Telugu.
"""

import io
import json
import os
from typing import Dict, List, Optional, Union

try:
    import docx
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    docx = None

try:
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
except ImportError:
    SimpleDocTemplate = None

from .engine import SUPPORTED_LANGUAGES


class DocumentExporter:
    """
    Exports multilingual translations into formatted DOCX, PDF, TXT, or JSON files/bytes.
    """

    # Registered font cache
    _FONTS_REGISTERED = False

    @classmethod
    def _init_reportlab_fonts(cls):
        """Register TTF fonts for Devanagari, Gujarati, Telugu, and Latin scripts in ReportLab."""
        if cls._FONTS_REGISTERED or not SimpleDocTemplate:
            return

        font_candidates = [
            ("Mangal", ["C:/Windows/Fonts/mangal.ttf", "C:/Windows/Fonts/Mangal.ttf"]),
            ("Shruti", ["C:/Windows/Fonts/shruti.ttf", "C:/Windows/Fonts/Shruti.ttf"]),
            ("Gautami", ["C:/Windows/Fonts/gautami.ttf", "C:/Windows/Fonts/Gautami.ttf"]),
            ("Arial", ["C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/Arial.ttf"]),
        ]

        for font_name, paths in font_candidates:
            for p in paths:
                if os.path.exists(p):
                    try:
                        pdfmetrics.registerFont(TTFont(font_name, p))
                        break
                    except Exception:
                        pass

        cls._FONTS_REGISTERED = True

    @classmethod
    def _get_font_for_lang(cls, lang_code: str) -> str:
        lang_code = lang_code.lower()
        if lang_code in ["hi", "mr"]:
            # Devanagari
            if "Mangal" in pdfmetrics.getRegisteredFontNames():
                return "Mangal"
        elif lang_code == "gu":
            # Gujarati
            if "Shruti" in pdfmetrics.getRegisteredFontNames():
                return "Shruti"
        elif lang_code == "te":
            # Telugu
            if "Gautami" in pdfmetrics.getRegisteredFontNames():
                return "Gautami"

        if "Arial" in pdfmetrics.getRegisteredFontNames():
            return "Arial"
        return "Helvetica"

    # ==========================
    # DOCX EXPORT
    # ==========================
    @classmethod
    def to_docx_bytes(
        cls,
        translated_text_or_chunks: Union[str, List[str]],
        lang_code: str,
        title: str = "Translated Document",
        original_title: str = "",
    ) -> bytes:
        if not docx:
            raise ImportError("python-docx is required for DOCX export.")

        doc = docx.Document()
        lang_info = SUPPORTED_LANGUAGES.get(lang_code, {"name": lang_code, "native": ""})
        font_name = lang_info.get("font_hint", "Arial")

        # Set default normal style font
        style = doc.styles["Normal"]
        font = style.font
        font.name = font_name
        font.size = Pt(11)
        font.color.rgb = RGBColor(30, 30, 30)

        # Header Title
        title_para = doc.add_paragraph()
        title_run = title_para.add_run(title)
        title_run.font.name = font_name
        title_run.font.size = Pt(18)
        title_run.bold = True
        title_run.font.color.rgb = RGBColor(26, 82, 118)

        # Subtitle / Language metadata
        meta_para = doc.add_paragraph()
        sub_run = meta_para.add_run(
            f"Language: {lang_info.get('name')} ({lang_info.get('native')}) | Target Code: [{lang_code}]"
        )
        sub_run.font.size = Pt(9.5)
        sub_run.italic = True
        sub_run.font.color.rgb = RGBColor(120, 120, 120)

        if original_title:
            sub2_run = meta_para.add_run(f" | Source: {original_title}")
            sub2_run.font.size = Pt(9.5)
            sub2_run.italic = True

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

        # Paragraphs
        if isinstance(translated_text_or_chunks, str):
            paragraphs = [p.strip() for p in translated_text_or_chunks.split("\n\n") if p.strip()]
        else:
            paragraphs = translated_text_or_chunks

        for p_text in paragraphs:
            p_text = p_text.strip()
            if not p_text:
                continue

            # Check if this paragraph looks like a heading
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.25

            if len(p_text) < 80 and not p_text.endswith((".", "।", "?", "!")):
                run = p.add_run(p_text)
                run.bold = True
                run.font.name = font_name
                run.font.size = Pt(13)
                run.font.color.rgb = RGBColor(40, 40, 40)
            else:
                run = p.add_run(p_text)
                run.font.name = font_name
                run.font.size = Pt(11)

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @classmethod
    def save_docx(
        cls,
        output_path: str,
        translated_text_or_chunks: Union[str, List[str]],
        lang_code: str,
        title: str = "Translated Document",
    ):
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        content = cls.to_docx_bytes(translated_text_or_chunks, lang_code, title)
        with open(output_path, "wb") as f:
            f.write(content)

    # ==========================
    # PDF EXPORT
    # ==========================
    @classmethod
    def to_pdf_bytes(
        cls,
        translated_text_or_chunks: Union[str, List[str]],
        lang_code: str,
        title: str = "Translated Document",
    ) -> bytes:
        if not SimpleDocTemplate:
            raise ImportError("reportlab is required for PDF export.")

        cls._init_reportlab_fonts()
        font_name = cls._get_font_for_lang(lang_code)

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            rightMargin=45,
            leftMargin=45,
            topMargin=45,
            bottomMargin=45,
        )

        lang_info = SUPPORTED_LANGUAGES.get(lang_code, {"name": lang_code, "native": ""})
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "DocTitle",
            fontName=font_name,
            fontSize=18,
            leading=22,
            textColor=colors.HexColor("#1A5276"),
            spaceAfter=6,
        )

        subtitle_style = ParagraphStyle(
            "DocSubTitle",
            fontName="Helvetica-Oblique",
            fontSize=9.5,
            leading=12,
            textColor=colors.HexColor("#7F8C8D"),
            spaceAfter=12,
        )

        body_style = ParagraphStyle(
            "DocBody",
            fontName=font_name,
            fontSize=11,
            leading=16,
            textColor=colors.HexColor("#2C3E50"),
            spaceAfter=10,
        )

        story = []
        story.append(Paragraph(cls._escape_xml(title), title_style))
        story.append(
            Paragraph(
                f"Target Language: {lang_info.get('name')} ({lang_info.get('native')}) | Code: [{lang_code}]",
                subtitle_style,
            )
        )
        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#BDC3C7"), spaceAfter=15))

        if isinstance(translated_text_or_chunks, str):
            paragraphs = [p.strip() for p in translated_text_or_chunks.split("\n\n") if p.strip()]
        else:
            paragraphs = translated_text_or_chunks

        for p_text in paragraphs:
            p_text = p_text.strip()
            if not p_text:
                continue

            escaped = cls._escape_xml(p_text).replace("\n", "<br/>")
            story.append(Paragraph(escaped, body_style))
            story.append(Spacer(1, 4))

        doc.build(story)
        buffer.seek(0)
        return buffer.getvalue()

    @classmethod
    def save_pdf(
        cls,
        output_path: str,
        translated_text_or_chunks: Union[str, List[str]],
        lang_code: str,
        title: str = "Translated Document",
    ):
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        content = cls.to_pdf_bytes(translated_text_or_chunks, lang_code, title)
        with open(output_path, "wb") as f:
            f.write(content)

    # ==========================
    # TXT & JSON EXPORT
    # ==========================
    @classmethod
    def to_txt_str(cls, translated_text_or_chunks: Union[str, List[str]]) -> str:
        if isinstance(translated_text_or_chunks, str):
            return translated_text_or_chunks
        return "\n\n".join(translated_text_or_chunks)

    @classmethod
    def save_txt(cls, output_path: str, translated_text_or_chunks: Union[str, List[str]]):
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        text = cls.to_txt_str(translated_text_or_chunks)
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(text)

    @classmethod
    def save_json(cls, output_path: str, data: Dict):
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        with open(output_path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)

    @staticmethod
    def _escape_xml(text: str) -> str:
        """Escapes XML entities for ReportLab Paragraphs."""
        return (
            text.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )

